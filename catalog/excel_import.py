"""Bounded XLSX parsing and tenant/branch-aware catalogue imports."""
from core.audit import audit_actor
from copy import copy
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from defusedxml.ElementTree import iterparse
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from core.models import Branch, Tenant, audit
from core.permissions import STOCK
from core.tenancy import check_staff_branch
from core.utils import decimal_input
from stock.models import StockMovement
from stock.services import record_movement
from .models import BranchStock, Category, CatalogueImport, Part, Supplier, Vehicle

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNPACKED_BYTES = 20 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 64
HEADERS = ['sku', 'name', 'category', 'brand', 'condition', 'oem_number', 'unit',
           'bin_location', 'cost_price', 'sell_price', 'trade_price', 'quantity',
           'reorder_level', 'barcode', 'part_type', 'reorder_qty', 'preferred_supplier',
           'description', 'is_universal', 'is_active', 'vehicle_make', 'vehicle_model',
           'year_from', 'year_to', 'engine']
ALIASES = {'part_code': 'sku', 'part_name': 'name', 'retail_price': 'sell_price',
           'opening_stock': 'quantity', 'supplier': 'preferred_supplier'}
MONEY = ('cost_price', 'sell_price', 'trade_price', 'quantity', 'reorder_level', 'reorder_qty')
TEXT_LIMITS = {'sku': 40, 'name': 160, 'category': 80, 'brand': 60, 'oem_number': 60,
               'unit': 20, 'bin_location': 40, 'barcode': 64, 'preferred_supplier': 120,
               'description': 2000, 'vehicle_make': 60, 'vehicle_model': 60, 'engine': 60}


class ImportProblem(ValueError):
    def __init__(self, errors):
        self.errors = errors[:100]
        self.total_errors = len(errors)
        super().__init__('; '.join(self.errors))


def parse_workbook(upload):
    if Path(upload.name).suffix.lower() != '.xlsx':
        raise ImportProblem(['Choose an Excel .xlsx file. Save older .xls files as .xlsx first.'])
    if upload.size > MAX_FILE_BYTES:
        raise ImportProblem(['The file exceeds the 5 MB upload limit.'])
    content = upload.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise ImportProblem(['The file exceeds the 5 MB upload limit.'])
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(m.file_size for m in members) > MAX_UNPACKED_BYTES:
                raise ImportProblem(['The workbook is too large when unpacked. Split it into smaller files.'])
            if any('vbaproject' in m.filename.lower() or 'externallinks/' in m.filename.lower() for m in members):
                raise ImportProblem(['Macros and links to external workbooks are not supported.'])
            # Check actual XML row/cell coordinates, not just workbook dimensions.
            for member in members:
                if member.filename.startswith('xl/worksheets/') and member.filename.endswith('.xml'):
                    with archive.open(member) as stream:
                        for event, element in iterparse(stream, events=('start', 'end')):
                            tag = element.tag.rsplit('}', 1)[-1]
                            if event == 'start' and tag == 'row' and int(element.get('r', '0')) > MAX_ROWS + 1:
                                raise ImportProblem(['Use no more than 5,000 catalogue rows plus the header. Remove extra formatted rows.'])
                            if event == 'start' and tag == 'mergeCell':
                                raise ImportProblem(['Merged cells are not supported. Unmerge cells before uploading.'])
                            if event == 'start' and tag == 'c':
                                address = element.get('r', '')
                                letters = ''.join(c for c in address if c.isalpha())
                                if letters and column_index_from_string(letters) > MAX_COLUMNS:
                                    raise ImportProblem(['The workbook exceeds the 64-column limit.'])
                            if event == 'end':
                                element.clear()
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        named = next((s for s in workbook.worksheets if s.title.lower() == 'catalogue'), None)
        sheet = named or next((s for s in workbook.worksheets if s.sheet_state == 'visible'), None)
        if sheet is None or sheet.sheet_state != 'visible':
            raise ImportProblem(['The workbook has no visible catalogue worksheet.'])
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(max_col=MAX_COLUMNS, max_row=MAX_ROWS + 1)
        first = next(iterator, ())
        columns = {}
        for index, cell in enumerate(first):
            if cell.value is None or str(cell.value).strip() == '':
                continue
            if cell.data_type == 'f' or not isinstance(cell.value, str):
                raise ImportProblem(['Row 1 must contain plain-text column headings.'])
            heading = cell.value.strip().lower().replace(' ', '_')
            heading = ALIASES.get(heading, heading)
            if heading not in HEADERS:
                raise ImportProblem([f'Unknown column "{cell.value}". Use the downloadable template headings.'])
            if heading in columns.values():
                raise ImportProblem([f'Duplicate column "{heading}".'])
            columns[index] = heading
        if not {'sku', 'name'}.issubset(columns.values()):
            raise ImportProblem(['Row 1 must include sku and name columns.'])
        rows, errors = [], []
        for number, cells in enumerate(iterator, start=2):
            if all(c.value is None or c.value == '' for c in cells):
                continue
            row = {'_row': number}
            for index, cell in enumerate(cells):
                value = cell.value
                if value is None or value == '':
                    continue
                if index not in columns:
                    errors.append(f'Row {number}: a value has no column heading.')
                    continue
                heading = columns[index]
                if cell.data_type in ('f', 'e'):
                    errors.append(f'Row {number}, {heading}: use a value, not a formula or Excel error.')
                    continue
                if heading in ('sku', 'barcode', 'oem_number') and not isinstance(value, str):
                    errors.append(f'Row {number}, {heading}: format identifiers as Text to preserve leading zeros.')
                    continue
                if not isinstance(value, (str, int, float, bool)):
                    errors.append(f'Row {number}, {heading}: unsupported cell value.')
                    continue
                row[heading] = str(value).strip()
            rows.append(row)
        if errors:
            raise ImportProblem(errors)
        if not rows:
            raise ImportProblem(['The catalogue sheet has no part rows. Fill in the template below its header.'])
        return rows
    except ImportProblem:
        raise
    except Exception as exc:
        # Parser failures are user-facing validation errors, never raw server errors.
        raise ImportProblem(['The workbook cannot be read. Upload an unencrypted .xlsx file using the template.']) from exc
    finally:
        if workbook is not None:
            workbook.close()


def validate_rows(rows, tenant, *, mode):
    errors, normalized, seen_skus, seen_barcodes = [], [], set(), set()
    skus = [(r.get('sku') or '').strip().upper() for r in rows]
    existing = {p.sku: p for p in Part.objects.filter(tenant=tenant, sku__in=skus)}
    barcodes = [r.get('barcode') for r in rows if r.get('barcode')]
    owners = dict(Part.objects.filter(tenant=tenant, barcode__in=barcodes).values_list('barcode', 'sku'))
    for raw in rows:
        number = raw['_row']; row = {'_row': number}
        try:
            sku = (raw.get('sku') or '').strip().upper()
            if not sku or not raw.get('name'):
                raise ValueError('sku and name are required for each part.')
            if sku in seen_skus:
                raise ValueError(f'duplicate SKU {sku} in this file. Use one row per part.')
            seen_skus.add(sku)
            part = existing.get(sku)
            if part and mode == 'create':
                raise ValueError(f'SKU {sku} already exists. Choose add and update mode if you intend to change it.')
            row['sku'] = sku
            for key, limit in TEXT_LIMITS.items():
                if raw.get(key):
                    if len(raw[key]) > limit:
                        raise ValueError(f'{key} exceeds {limit} characters.')
                    row[key] = raw[key]
            row['sku'] = sku
            if row.get('barcode'):
                code = row['barcode']
                if code in seen_barcodes or (code in owners and owners[code] != sku):
                    raise ValueError(f'barcode {code} is already assigned to another part.')
                seen_barcodes.add(code)
            for key in MONEY:
                if raw.get(key) not in (None, ''):
                    row[key] = str(decimal_input(raw[key], label=key, minimum=0))
            if part is None and 'sell_price' not in row:
                raise ValueError('sell_price is required for a new part. Enter 0 explicitly for a free item.')
            for key, choices in [('condition', Part.Condition.values), ('part_type', Part.PartType.values)]:
                if raw.get(key):
                    value = raw[key].lower()
                    if value not in choices:
                        raise ValueError(f'{key} must be one of: {", ".join(choices)}.')
                    row[key] = value
            for key in ('is_universal', 'is_active'):
                if raw.get(key) not in (None, ''):
                    value = str(raw[key]).lower()
                    if value not in ('true', 'false', 'yes', 'no', '1', '0'):
                        raise ValueError(f'{key} must be yes/no or true/false.')
                    row[key] = value in ('true', 'yes', '1')
            if any(raw.get(k) for k in ('vehicle_make', 'vehicle_model', 'year_from', 'year_to', 'engine')):
                if not all(raw.get(k) for k in ('vehicle_make', 'vehicle_model', 'year_from', 'year_to')):
                    raise ValueError('vehicle_make, vehicle_model, year_from and year_to must be provided together.')
                for key in ('year_from', 'year_to'):
                    value = Decimal(raw[key])
                    if not value.is_finite() or value != int(value) or not 1886 <= value <= 2100:
                        raise ValueError(f'{key} must be a whole year from 1886 to 2100.')
                    row[key] = int(value)
                if row['year_to'] < row['year_from']:
                    raise ValueError('year_to cannot be before year_from.')
            candidate = copy(part) if part else Part(tenant=tenant, sku=sku)
            for key in PART_FIELDS:
                if key in row:
                    setattr(candidate, key, Decimal(row[key]) if key in MONEY else row[key])
            candidate.full_clean(exclude=['tenant', 'category', 'preferred_supplier'], validate_unique=False, validate_constraints=False)
            normalized.append(row)
        except (ValueError, ValidationError, ArithmeticError) as exc:
            message = '; '.join(exc.messages) if isinstance(exc, ValidationError) else str(exc)
            errors.append(f'Row {number}: {message}')
    if errors:
        raise ImportProblem(errors)
    return normalized


PART_FIELDS = ('name', 'brand', 'condition', 'oem_number', 'unit', 'cost_price',
               'sell_price', 'trade_price', 'barcode', 'part_type', 'description', 'is_universal', 'is_active')


def import_plan(rows, tenant, branch):
    existing = {p.sku: p for p in Part.objects.filter(tenant=tenant, sku__in=[r['sku'] for r in rows])}
    stocks = {s.part_id: s for s in BranchStock.objects.filter(branch=branch, part__in=existing.values())}
    histories = set(StockMovement.objects.filter(branch=branch, part__in=existing.values()).values_list('part_id', flat=True))
    plan = []
    for row in rows:
        part = existing.get(row['sku']); stock = stocks.get(part.pk) if part else None
        opening = Decimal(row.get('quantity', '0'))
        can_open = not part or (part.pk not in histories and (stock is None or stock.quantity_on_hand == 0))
        plan.append(dict(row, row_number=row['_row'], action='Update' if part else 'Add', price=row.get('sell_price', str(part.sell_price) if part else '0'),
                         opening=str(opening) if can_open else '0', ignored_opening=opening > 0 and not can_open))
    return plan


@audit_actor
@transaction.atomic
def confirm_import(batch, user):
    if user.role not in STOCK or batch.created_by_id != user.pk or batch.tenant_id != user.tenant_id:
        raise PermissionDenied('This upload belongs to another account.')
    check_staff_branch(user, batch.branch)
    batch = CatalogueImport.objects.select_for_update().get(pk=batch.pk)
    if batch.status == CatalogueImport.Status.IMPORTED:
        return batch.result
    if batch.status != CatalogueImport.Status.READY or timezone.now() > batch.created_at + timedelta(hours=1):
        raise ImportProblem(['This preview has expired or was cancelled. Upload the file again.'])
    Tenant.objects.select_for_update().get(pk=batch.tenant_id)
    rows = validate_rows(batch.rows, batch.tenant, mode=batch.mode)
    result = {'created': 0, 'updated': 0, 'openings': 0, 'ignored_openings': 0}
    branches = list(Branch.objects.filter(tenant=batch.tenant))
    for row in rows:
        part = Part.objects.filter(tenant=batch.tenant, sku=row['sku']).first()
        created = part is None
        part = part or Part(tenant=batch.tenant, sku=row['sku'])
        for key in PART_FIELDS:
            if key in row:
                setattr(part, key, Decimal(row[key]) if key in MONEY else row[key])
        for key, model in [('category', Category), ('preferred_supplier', Supplier)]:
            if row.get(key):
                setattr(part, key, model.objects.get_or_create(tenant=batch.tenant, name=row[key])[0])
        part.full_clean()
        part.save()
        BranchStock.objects.bulk_create([BranchStock(branch=b, part=part, cost_price=part.cost_price) for b in branches], ignore_conflicts=True)
        item = BranchStock.objects.select_for_update().get(branch=batch.branch, part=part)
        for key in ('bin_location', 'reorder_level', 'reorder_qty'):
            if key in row:
                setattr(item, key, Decimal(row[key]) if key != 'bin_location' else row[key])
        item.full_clean()
        item.save(update_fields=['bin_location', 'reorder_level', 'reorder_qty'])
        quantity = Decimal(row.get('quantity', '0'))
        has_history = StockMovement.objects.filter(branch=batch.branch, part=part).exists()
        if quantity > 0 and item.quantity_on_hand == 0 and not has_history:
            item.cost_price = part.cost_price
            item.save(update_fields=['cost_price'])
            record_movement(part, quantity, 'opening', user, reference=f'IMPORT-{batch.pk}', note='Excel catalogue opening stock', branch=batch.branch)
            result['openings'] += 1
        elif quantity > 0:
            result['ignored_openings'] += 1
        if row.get('vehicle_make'):
            vehicle, _ = Vehicle.objects.get_or_create(tenant=batch.tenant, make=row['vehicle_make'], model=row['vehicle_model'],
                year_from=row['year_from'], year_to=row['year_to'], engine=row.get('engine', ''))
            part.fits.add(vehicle)
        result['created' if created else 'updated'] += 1
    batch.status = CatalogueImport.Status.IMPORTED
    batch.completed_at = timezone.now()
    batch.result = result
    batch.save(update_fields=['status', 'completed_at', 'result'])
    audit(user, 'catalog.import', f'IMPORT-{batch.pk}', f"Excel: {result['created']} added, {result['updated']} updated, {result['openings']} branch openings", branch=batch.branch)
    return result
